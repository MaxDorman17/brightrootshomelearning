"""A parent can give any planned lesson a score, whatever scheme it comes from."""
from datetime import date, timedelta

from conftest import sign_up


def planned(family, child_id, title="Fractions", subject="Maths", scheme="White Rose Maths"):
    lesson = family.parent.post("/api/lessons/", json={"title": title, "subject": subject, "scheme": scheme}).json()
    body = {"lesson_id": lesson["id"], "scheduled_date": date.today().isoformat()}
    if child_id is not None:
        body["assigned_to"] = child_id
    entry = family.parent.post("/api/planner/", json=body)
    assert entry.status_code == 201, entry.text
    return entry.json()


def test_score_a_lesson_and_change_it(family):
    child = family.add_child()
    entry = planned(family, child["id"])
    made = family.parent.put(f"/api/test-results/lesson/{entry['id']}", json={"child_id": child["id"], "score": 8, "total": 10})
    assert made.status_code == 200, made.text
    assert made.json()["title"] == "Fractions"
    assert made.json()["subject"] == "Maths"
    assert made.json()["taken_on"] == date.today().isoformat()

    # Scoring it again changes the same result rather than adding another.
    family.parent.put(f"/api/test-results/lesson/{entry['id']}", json={"child_id": child["id"], "score": 9, "total": 10})
    scores = family.parent.get("/api/test-results/lesson-scores").json()
    assert [(s["entry_id"], s["child_id"], s["score"], s["total"]) for s in scores] == [(entry["id"], child["id"], 9, 10)]

    # It sits with the family's other results, and in the council report.
    overview = family.parent.get("/api/test-results/overview", params={"child_id": child["id"]}).json()
    assert [(t["title"], t["score"], t["entry_id"]) for t in overview["tests"]] == [("Fractions", 9, entry["id"])]
    today = date.today().isoformat()
    report = family.parent.get("/api/council-report/", params={"child_id": child["id"], "start_date": today, "end_date": today}).json()
    assert [(t["title"], t["percent"]) for t in report["results"]["tests"]] == [("Fractions", 90)]

    # The child can see their own score but can't set one.
    kid = family.child_client(child)
    assert [s["score"] for s in kid.get("/api/test-results/lesson-scores").json()] == [9]
    assert kid.put(f"/api/test-results/lesson/{entry['id']}", json={"child_id": child["id"], "score": 10, "total": 10}).status_code == 403

    # And it can be taken off again.
    assert family.parent.delete(f"/api/test-results/lesson/{entry['id']}", params={"child_id": child["id"]}).status_code == 204
    assert family.parent.get("/api/test-results/lesson-scores").json() == []


def test_scores_are_checked(family):
    child = family.add_child()
    entry = planned(family, child["id"])
    url = f"/api/test-results/lesson/{entry['id']}"
    assert family.parent.put(url, json={"child_id": child["id"], "score": 11, "total": 10}).status_code == 400
    assert family.parent.put(url, json={"child_id": child["id"], "score": 5, "total": 0}).status_code == 422
    assert family.parent.put(url, json={"child_id": child["id"], "score": -1, "total": 10}).status_code == 422
    # Not a lesson that child does.
    sibling = family.add_child("Sib")
    assert family.parent.put(url, json={"child_id": sibling["id"], "score": 5, "total": 10}).status_code == 404


def test_a_shared_lesson_has_a_score_for_each_child(family):
    one, two = family.add_child("One"), family.add_child("Two")
    entry = planned(family, None)
    url = f"/api/test-results/lesson/{entry['id']}"
    assert family.parent.put(url, json={"child_id": one["id"], "score": 7, "total": 10}).status_code == 200
    assert family.parent.put(url, json={"child_id": two["id"], "score": 4, "total": 10}).status_code == 200
    scores = {s["child_id"]: s["score"] for s in family.parent.get("/api/test-results/lesson-scores").json()}
    assert scores == {one["id"]: 7, two["id"]: 4}


def test_another_family_cannot_score_or_see_it(family):
    child = family.add_child()
    entry = planned(family, child["id"])
    family.parent.put(f"/api/test-results/lesson/{entry['id']}", json={"child_id": child["id"], "score": 8, "total": 10})
    other = sign_up("Other")
    their_child = other.add_child()
    assert other.parent.put(f"/api/test-results/lesson/{entry['id']}", json={"child_id": their_child["id"], "score": 1, "total": 10}).status_code == 404
    assert other.parent.put(f"/api/test-results/lesson/{entry['id']}", json={"child_id": child["id"], "score": 1, "total": 10}).status_code == 404
    assert other.parent.get("/api/test-results/lesson-scores").json() == []
    assert other.parent.delete(f"/api/test-results/lesson/{entry['id']}", params={"child_id": child["id"]}).status_code == 204
    assert len(family.parent.get("/api/test-results/lesson-scores").json()) == 1


def test_removing_the_lesson_keeps_the_result(family):
    child = family.add_child()
    entry = planned(family, child["id"])
    family.parent.put(f"/api/test-results/lesson/{entry['id']}", json={"child_id": child["id"], "score": 8, "total": 10})
    assert family.parent.delete(f"/api/planner/{entry['id']}").status_code == 204
    assert family.parent.get("/api/test-results/lesson-scores").json() == []
    overview = family.parent.get("/api/test-results/overview", params={"child_id": child["id"]}).json()
    assert [(t["title"], t["entry_id"]) for t in overview["tests"]] == [("Fractions", None)]


def test_a_good_score_earns_stars_and_counts_in_a_challenge(family):
    child = family.add_child()
    kid = family.child_client(child)
    setup = family.parent.get("/api/rewards/setup").json()  # opening Rewards adds the example rules
    assert any(r["kind"] == "score" and r["threshold_pct"] == 80 for r in setup["rules"])
    today = date.today()
    challenge = family.parent.post("/api/family/challenges", json={
        "title": "Three good scores", "kind": "score", "target": 3, "threshold_pct": 70,
        "start_date": today.isoformat(), "end_date": (today + timedelta(days=7)).isoformat(),
    })
    assert challenge.status_code == 201, challenge.text

    good = planned(family, child["id"], title="Fractions")
    poor = planned(family, child["id"], title="Decimals", subject="Science")
    family.parent.put(f"/api/test-results/lesson/{good['id']}", json={"child_id": child["id"], "score": 9, "total": 10})
    family.parent.put(f"/api/test-results/lesson/{poor['id']}", json={"child_id": child["id"], "score": 3, "total": 10})

    assert kid.get("/api/rewards/me").json()["available"] == 3  # 90% beats the 80% rule; 30% doesn't
    history = family.parent.get("/api/rewards/history", params={"child_id": child["id"]}).json()
    assert "Scored 90%: Fractions" in str(history)
    board = family.parent.get("/api/family/overview").json()
    assert "'progress': 1" in str(board) or '"progress": 1' in str(board)


def test_the_results_export_includes_marked_scores(family):
    import io

    from openpyxl import load_workbook

    child = family.add_child("Robin")
    entry = planned(family, child["id"])
    family.parent.put(f"/api/test-results/lesson/{entry['id']}", json={"child_id": child["id"], "score": 8, "total": 10})
    family.parent.post("/api/test-results/", json={
        "child_id": child["id"], "subject": "Maths", "title": "End of term paper",
        "taken_on": date.today().isoformat(), "score": 30, "total": 40,
    })
    other = sign_up("Other")
    their_child = other.add_child()
    their_entry = planned(other, their_child["id"], title="Not ours")
    other.parent.put(f"/api/test-results/lesson/{their_entry['id']}", json={"child_id": their_child["id"], "score": 1, "total": 10})

    exported = family.parent.get("/api/oak/export", params={"child_id": child["id"]})
    assert exported.status_code == 200, exported.text
    sheet = load_workbook(io.BytesIO(exported.content))["Lesson scores and tests"]
    rows = [[c.value for c in row] for row in sheet.iter_rows(min_row=2)]
    assert [(r[1], r[3], r[4], r[5], r[6]) for r in rows] == [
        ("Robin", "Fractions", "Lesson score", 8, 10),
        ("Robin", "End of term paper", "Test", 30, 40),
    ]
